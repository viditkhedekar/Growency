from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path('/Users/Vidit/Projects/Growency/outputs')
text = (ROOT / 'infratex-direct-b2b-appendix.md').read_text()
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin, s.bottom_margin = Inches(.65), Inches(.65)
s.left_margin, s.right_margin = Inches(.75), Inches(.75)
s.footer_distance = Inches(.28)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.widow_control = True
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.0
doc.styles['Normal'].paragraph_format.space_after = Pt(3)
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].paragraph_format.space_after = Pt(14)
doc.styles['Heading 1'].font.size = Pt(19)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].font.bold = True
doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)
doc.core_properties.title = 'Infratex direct B2B sales appendix'
doc.core_properties.subject = 'Direct B2B workflow offers and SDR execution'
doc.core_properties.author = 'Growency'

def hyperlink(p, label, url):
    rel = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rel)
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '195B8A'); props.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
    run.append(props); t = OxmlElement('w:t'); t.text = label; run.append(t)
    h.append(run); p._p.append(h)

def table(lines):
    rows = [[x.strip() for x in line.strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r':?-+:?',x) for x in r)]
    tab = doc.add_table(rows=0, cols=len(rows[0]))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    widths = [1.6, 3.7, 1.7]
    for c,w in zip(tab.columns,widths): c.width = Inches(w)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    tab._tbl.tblPr.append(borders)
    for n,vals in enumerate(rows):
        row = tab.add_row()
        if n == 0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for j,(cell,val) in enumerate(zip(row.cells,vals)):
            cell.width = Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'20384C' if n==0 else ('F3F6F8' if n%2==0 else 'FFFFFF'));pr.append(shade)
            mar=OxmlElement('w:tcMar')
            for e,v in [('top','95'),('bottom','95'),('left','110'),('right','110')]:
                item=OxmlElement('w:'+e);item.set(qn('w:w'),v);item.set(qn('w:type'),'dxa');mar.append(item)
            pr.append(mar)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.03
            run=p.add_run(val);run.font.size=Pt(10.5);run.bold=(n==0)
            if n==0:run.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

lines=text.splitlines();i=0;title=True;new_page=False
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line=='<!-- page -->':new_page=True;i+=1;continue
    if line.startswith('|'):
        group=[]
        while i<len(lines) and lines[i].strip().startswith('|'):group.append(lines[i].strip());i+=1
        table(group);continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],style='Title' if title else 'Heading 1')
        if new_page:p.paragraph_format.page_break_before=True;new_page=False
        title=False
    elif line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 2')
    elif line.startswith('- '):
        p=doc.add_paragraph(line[2:],style='List Bullet');p.paragraph_format.space_after=Pt(3)
    elif re.match(r'^\[S\d+\].+https://',line):
        label,url=line.rsplit(' — ',1);p=doc.add_paragraph();p.add_run(label+'  ');hyperlink(p,'Open source',url)
        p.paragraph_format.space_after=Pt(6)
    else:doc.add_paragraph(line)
    i+=1

p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=p.add_run('Growency   |   Direct B2B appendix   |   ');r.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
output=ROOT/'infratex-direct-b2b-appendix.docx';doc.save(output)
print(output)
print('Source words:',len(text.split()))
print('Planned pages:',text.count('<!-- page -->')+1)
